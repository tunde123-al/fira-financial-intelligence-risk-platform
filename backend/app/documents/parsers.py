"""Document parsers: Markdown/text, PDF (with OCR fallback), DOCX, and Docling.

Docling is the preferred parser for complex layouts when installed
(`pip install docling`, see requirements-docling.txt) and enabled with
DOCUMENT_PARSER=docling. It is heavy (it pulls ML layout models), so the default
build uses the lightweight parsers below, which handle text PDFs, scanned PDFs
(Tesseract OCR), DOCX headings/tables and Markdown.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any

import yaml

from app.documents.models import Block

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+([A-Z][^.!?]{2,80})$")
FRONT_MATTER_RE = re.compile(r"^---\s*\n(.*?)\n---\s*\n", re.S)


class ParseResult:
    def __init__(self, blocks: list[Block], metadata: dict[str, Any], page_count: int | None, parser: str,
                 ocr_pages: list[int] | None = None):
        self.blocks = blocks
        self.metadata = metadata
        self.page_count = page_count
        self.parser = parser
        self.ocr_pages = ocr_pages or []


def _assign_sections(blocks: list[Block]) -> list[Block]:
    path: list[tuple[int, str]] = []
    for b in blocks:
        if b.kind == "heading":
            path = [p for p in path if p[0] < b.level] + [(b.level, b.text)]
            b.section_path = [p[1] for p in path]
        else:
            b.section_path = [p[1] for p in path]
    return blocks


def parse_markdown(text: str) -> ParseResult:
    meta: dict[str, Any] = {}
    m = FRONT_MATTER_RE.match(text)
    if m:
        meta = yaml.safe_load(m.group(1)) or {}
        text = text[m.end():]
    blocks: list[Block] = []
    para: list[str] = []
    kind: Any = "paragraph"

    def flush() -> None:
        nonlocal para, kind
        if para:
            blocks.append(Block(kind=kind, text="\n".join(para).strip()))
        para, kind = [], "paragraph"

    for line in text.splitlines():
        h = HEADING_RE.match(line)
        if h:
            flush()
            blocks.append(Block(kind="heading", text=h.group(2).strip(), level=len(h.group(1))))
            continue
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("|"):
            new_kind = "table"
        elif re.match(r"^([-*]|\d+\.)\s+", stripped):
            new_kind = "list"
        elif stripped.startswith(">"):
            new_kind = "quote"
        elif kind == "list" and para:
            new_kind = "list"  # wrapped continuation of a list item
        else:
            new_kind = "paragraph"
        if para and new_kind != kind:
            flush()
        kind = new_kind
        para.append(stripped.lstrip(">").strip() if new_kind == "quote" else stripped)
    flush()
    return ParseResult(_assign_sections(blocks), meta, None, "markdown")


def parse_text(text: str) -> ParseResult:
    blocks = []
    for para in re.split(r"\n\s*\n", text):
        p = para.strip()
        if not p:
            continue
        nh = NUMBERED_HEADING_RE.match(p) if "\n" not in p else None
        if nh:
            blocks.append(Block(kind="heading", text=p, level=nh.group(1).count(".") + 1))
        else:
            blocks.append(Block(kind="paragraph", text=p))
    return ParseResult(_assign_sections(blocks), {}, None, "text")


def _table_to_text(rows: list[list[Any]]) -> str:
    clean = [[("" if c is None else str(c).replace("\n", " ").strip()) for c in r] for r in rows if r]
    if not clean:
        return ""
    lines = ["| " + " | ".join(clean[0]) + " |", "|" + "---|" * len(clean[0])]
    lines += ["| " + " | ".join(r) + " |" for r in clean[1:]]
    return "\n".join(lines)


def _ocr_page(page: Any) -> str:
    import pytesseract

    img = page.to_image(resolution=200).original
    return pytesseract.image_to_string(img)


def _blocks_from_page_text(text: str, page_no: int) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            if para:
                blocks.append(Block(kind="paragraph", text=" ".join(para), page=page_no))
                para = []
            continue
        nh = NUMBERED_HEADING_RE.match(s)
        if nh or (s.isupper() and 3 < len(s) < 80):
            if (not nh and blocks and blocks[-1].kind == "heading" and not para
                    and blocks[-1].text.isupper()):
                blocks[-1].text += " " + s  # heading wrapped over two lines
                continue
            if para:
                blocks.append(Block(kind="paragraph", text=" ".join(para), page=page_no))
                para = []
            level = nh.group(1).count(".") + 1 if nh else 1
            blocks.append(Block(kind="heading", text=s, level=level, page=page_no))
        else:
            para.append(s)
    if para:
        blocks.append(Block(kind="paragraph", text=" ".join(para), page=page_no))
    return blocks


def parse_pdf(path: Path, ocr: bool = True, min_chars_for_text_page: int = 25) -> ParseResult:
    import pdfplumber

    blocks: list[Block] = []
    ocr_pages: list[int] = []
    meta: dict[str, Any] = {}
    with pdfplumber.open(str(path)) as pdf:
        meta = {k.lower(): v for k, v in (pdf.metadata or {}).items() if isinstance(v, str)}
        for i, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []
            text = page.extract_text() or ""
            if len(text.strip()) < min_chars_for_text_page and ocr:
                try:
                    text = _ocr_page(page)
                    ocr_pages.append(i)
                except Exception as e:  # OCR unavailable: keep whatever text exists
                    meta.setdefault("ocr_errors", []).append(f"page {i}: {type(e).__name__}")
            if tables:
                # remove table text lines from the page body to avoid duplicates
                table_cells = {str(c).strip() for t in tables for r in t for c in r if c}
                table_rows = {" ".join(str(c).strip() for c in r if c) for t in tables for r in t}
                text = "\n".join(ln for ln in text.splitlines()
                                  if ln.strip() not in table_cells and ln.strip() not in table_rows)
            blocks.extend(_blocks_from_page_text(text, i))
            for t in tables:
                tt = _table_to_text(t)
                if tt:
                    blocks.append(Block(kind="table", text=tt, page=i))
        page_count = len(pdf.pages)
    return ParseResult(_assign_sections(blocks), meta, page_count, "pdfplumber+ocr" if ocr_pages else "pdfplumber",
                       ocr_pages)


def parse_docx(path: Path) -> ParseResult:
    import docx

    d = docx.Document(str(path))
    blocks: list[Block] = []
    body = d.element.body
    paras = {p._p: p for p in d.paragraphs}
    tables = {t._tbl: t for t in d.tables}
    for el in body.iterchildren():
        if el in paras:
            p = paras[el]
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.match(r"heading (\d)", style)
            if m or style == "title":
                blocks.append(Block(kind="heading", text=text, level=int(m.group(1)) if m else 1))
            elif "list" in style:
                blocks.append(Block(kind="list", text=text))
            else:
                blocks.append(Block(kind="paragraph", text=text))
        elif el in tables:
            t = tables[el]
            rows = [[c.text for c in r.cells] for r in t.rows]
            tt = _table_to_text(rows)
            if tt:
                blocks.append(Block(kind="table", text=tt))
    cp = d.core_properties
    meta = {k: v for k, v in {"title": cp.title, "author": cp.author, "version": cp.version}.items() if v}
    # merge consecutive list items
    merged: list[Block] = []
    for b in blocks:
        if merged and b.kind == "list" and merged[-1].kind == "list":
            merged[-1].text += "\n" + b.text
        else:
            merged.append(b)
    return ParseResult(_assign_sections(merged), meta, None, "python-docx")


def parse_with_docling(path: Path) -> ParseResult:  # pragma: no cover - optional heavy dependency
    from docling.document_converter import DocumentConverter

    result = DocumentConverter().convert(str(path))
    md = result.document.export_to_markdown()
    pr = parse_markdown(md)
    pr.parser = "docling"
    pages = getattr(result.document, "pages", None)
    pr.page_count = len(pages) if pages is not None else None
    return pr


def parse_file(path: Path) -> ParseResult:
    path = Path(path)
    suffix = path.suffix.lower()
    use_docling = os.environ.get("DOCUMENT_PARSER", "").lower() == "docling"
    if use_docling and suffix in {".pdf", ".docx"}:
        try:
            return parse_with_docling(path)
        except ImportError:
            pass
    if suffix in {".md", ".markdown"}:
        return parse_markdown(path.read_text(encoding="utf-8"))
    if suffix == ".txt":
        return parse_text(path.read_text(encoding="utf-8"))
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"unsupported document type: {suffix}")


SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".docx"}
