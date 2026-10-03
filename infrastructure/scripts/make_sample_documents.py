"""Generate the binary sample documents (PDF with a table, scanned image-only PDF,
DOCX contract) used to exercise the PDF/OCR/DOCX ingestion paths.

All content is synthetic and refers to the fictional Lagoon Bank Plc.
Usage: python infrastructure/scripts/make_sample_documents.py
"""
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2] / "documents"


def make_pdf_report(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title="Financial Crime MI Report Q2 2026",
                            author="Lagoon Bank Plc (fictional)")
    story = [
        Paragraph("FINANCIAL CRIME MANAGEMENT INFORMATION REPORT Q2 2026", ss["Title"]),
        Paragraph("Synthetic document for FIRA development. Lagoon Bank Plc is a fictional institution. "
                  "Version 1.0. Effective date 2026-07-15.", ss["Italic"]),
        Spacer(1, 12),
        Paragraph("1 Executive Summary", ss["Heading2"]),
        Paragraph("Alert volumes increased in the second quarter, driven mainly by card testing bursts on online "
                  "merchants and by money mule activity involving rapid pass-through of funds. The false-positive "
                  "rate of the legacy velocity rule remained high because it does not compare activity with the "
                  "customer's own baseline.", ss["BodyText"]),
        Paragraph("2 Alert Volumes by Scenario", ss["Heading2"]),
        Table([["Scenario", "Alerts", "Escalated", "False positive rate"],
               ["Legacy velocity", "1,240", "38", "91%"],
               ["Large transfer", "860", "41", "84%"],
               ["Geographic", "415", "22", "88%"],
               ["Mule pass-through", "212", "57", "61%"]],
              style=TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)])),
        PageBreak(),
        Paragraph("3 Observations", ss["Heading2"]),
        Paragraph("Device sharing among unrelated customers was the strongest single network indicator in confirmed "
                  "mule cases. Dormant accounts reactivated with large deposits followed by transfers out within a "
                  "day were frequently linked to accounts sold to third parties.", ss["BodyText"]),
        Paragraph("4 Recommendations", ss["Heading2"]),
        Paragraph("Replace the legacy velocity rule with baseline-relative burst detection; add graph-based "
                  "detection of circular flows; require investigators to document alternative legitimate "
                  "explanations.", ss["BodyText"]),
    ]
    doc.build(story)


def make_scanned_pdf(path: Path) -> None:
    import img2pdf
    from PIL import Image, ImageDraw, ImageFont

    lines = [
        "LAGOON BANK PLC (FICTIONAL) - INTERNAL MEMO",
        "",
        "Subject: Interim guidance on card testing",
        "Date: 2026-08-20",
        "",
        "Fraud Operations has observed bursts of small online card",
        "payments with many declined authorisations. Analysts should",
        "treat twenty or more card-not-present attempts within one",
        "hour as a card testing indicator and review the card for",
        "compromise. Blocking the card requires approval by an",
        "authorised Fraud Operations officer.",
        "",
        "This memo is a synthetic document for FIRA development.",
    ]
    img = Image.new("L", (1700, 2200), 255)
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 40)
    except OSError:
        font = ImageFont.load_default(size=40)
    y = 150
    for line in lines:
        draw.text((120, y), line, fill=0, font=font)
        y += 70
    png = path.with_suffix(".png")
    img.save(png, dpi=(200, 200))
    path.write_bytes(img2pdf.convert(str(png)))
    png.unlink()


def make_docx_contract(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.title = "Correspondent Banking Services Agreement (Template)"
    d.core_properties.author = "Lagoon Bank Plc (fictional)"
    d.add_heading("Correspondent Banking Services Agreement (Template)", 0)
    d.add_paragraph("Synthetic contract template for FIRA development. Lagoon Bank Plc is fictional. Version 0.9.")
    d.add_heading("1. Parties", 1)
    d.add_paragraph("This agreement is made between Lagoon Bank Plc (the Correspondent) and the Respondent "
                    "institution named in Schedule 1. The parties hereby agree to the terms below.")
    d.add_heading("2. Financial Crime Obligations", 1)
    d.add_paragraph("The Respondent shall maintain an anti-money laundering programme including customer due "
                    "diligence, transaction monitoring and suspicious activity escalation.", style="List Number")
    d.add_paragraph("The Respondent shall respond to requests for information about specific transactions within "
                    "five business days.", style="List Number")
    d.add_paragraph("The Correspondent may suspend services where the Respondent fails to meet these obligations, "
                    "subject to the notice provisions in clause 4.", style="List Number")
    d.add_heading("3. Service Fees", 1)
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Service", "Fee"), ("Payment processing", "USD 4.00 per item"),
                                ("Request for information handling", "No charge")]):
        t.rows[i].cells[0].text, t.rows[i].cells[1].text = a, b
    d.add_heading("4. Termination", 1)
    d.add_paragraph("Either party may terminate this agreement with ninety days written notice. Governing law is "
                    "specified in Schedule 2.")
    d.save(str(path))


def main() -> None:
    (ROOT / "reports").mkdir(parents=True, exist_ok=True)
    (ROOT / "contracts").mkdir(parents=True, exist_ok=True)
    (ROOT / "memos").mkdir(parents=True, exist_ok=True)
    make_pdf_report(ROOT / "reports" / "fincrime_mi_report_q2_2026.pdf")
    make_scanned_pdf(ROOT / "memos" / "card_testing_memo_scanned.pdf")
    make_docx_contract(ROOT / "contracts" / "correspondent_banking_agreement_template.docx")
    print("sample documents written")


if __name__ == "__main__":
    main()
